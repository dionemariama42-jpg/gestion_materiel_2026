from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone
from datetime import timedelta
from emprunts.models import Demande, Restitution
from materiel.models import Materiel, Categorie, Maintenance
from comptes.models import Utilisateur
from clubs.models import Club, Activite, MembreClub
from cahier.models import Cours, Seance, FichierSeance
import json


@login_required
def dashboard_admin(request):
    aujourd_hui = timezone.now().date()
    debut_mois = aujourd_hui.replace(day=1)
    debut_annee = aujourd_hui.replace(month=1, day=1)

    from django.db.models import Count, Q

    # ── Matériels ──
    total_materiels = Materiel.objects.count()
    materiels_disponibles = Materiel.objects.filter(etat='disponible').count()
    materiels_empruntes = Materiel.objects.filter(etat='emprunte').count()
    materiels_maintenance = Materiel.objects.filter(etat='maintenance').count()
    materiels_hors_service = Materiel.objects.filter(etat='hors_service').count()
    taux_utilisation = round((materiels_empruntes / total_materiels * 100) if total_materiels else 0)

    # ── Demandes ──
    total_demandes = Demande.objects.count()
    demandes_en_attente = Demande.objects.filter(statut='en_attente').count()
    demandes_en_cours = Demande.objects.filter(statut='en_cours').count()
    demandes_restituees = Demande.objects.filter(statut='restituee').count()
    demandes_refusees = Demande.objects.filter(statut='refusee').count()
    demandes_ce_mois = Demande.objects.filter(date_demande__date__gte=debut_mois).count()
    demandes_cette_annee = Demande.objects.filter(date_demande__date__gte=debut_annee).count()
    retards = Demande.objects.filter(statut='en_cours', date_fin__lt=aujourd_hui).count()

    # ── Utilisateurs ──
    total_etudiants = Utilisateur.objects.filter(role='etudiant').count()
    total_enseignants = Utilisateur.objects.filter(role='enseignant').count()
    total_techniciens = Utilisateur.objects.filter(role='technicien').count()
    etudiants_bloques = Utilisateur.objects.filter(bloque=True).count()
    total_penalites = sum(u.penalite for u in Utilisateur.objects.all())

    # ── Clubs ──
    total_clubs = Club.objects.count()
    total_membres = MembreClub.objects.count()
    total_activites = Activite.objects.count()
    activites_ce_mois = Activite.objects.filter(date__gte=debut_mois).count()
    activites_terminees = Activite.objects.filter(statut='terminee').count()
    # ── Cahier de textes ──
    total_cours = Cours.objects.count()
    total_seances = Seance.objects.count()
    seances_ce_mois = Seance.objects.filter(date__gte=debut_mois).count()
    total_fichiers = FichierSeance.objects.count()
    top_enseignants = Utilisateur.objects.filter(
    role='enseignant'
    ).annotate(
    nb_seances=Count('seances')
    ).order_by('-nb_seances')[:5]
    dernieres_seances = Seance.objects.all().order_by('-date')[:5]

    # ── Maintenances ──
    maintenances_en_cours = Maintenance.objects.filter(statut='en_cours').count()
    maintenances_en_attente = Maintenance.objects.filter(statut='en_attente').count()

    # ── Dernières demandes ──
    dernieres_demandes = Demande.objects.all().order_by('-date_demande')[:10]

    # ── Top matériels ──
    top_materiels = Materiel.objects.annotate(
        nb_emprunts=Count('lignes_demande')
    ).order_by('-nb_emprunts')[:8]

    # ── Matériels par catégorie ──
    categories_stats = Categorie.objects.annotate(
        nb_materiels=Count('materiels')
    ).order_by('-nb_materiels')

    # ── Clubs stats ──
    clubs_stats = Club.objects.annotate(
        nb_membres=Count('membres'),
        nb_activites=Count('activites')
    ).order_by('-nb_activites')

    # ── Top pénalités ──
    top_penalites = Utilisateur.objects.filter(
        role='etudiant', penalite__gt=0
    ).order_by('-penalite')[:5]

    # ── Étudiants les plus actifs ──
    top_etudiants = Utilisateur.objects.filter(
        role='etudiant'
    ).annotate(
        nb_emprunts=Count('demandes')
    ).order_by('-nb_emprunts')[:5]

    # ── Graphique emprunts par mois (12 mois) ──
    mois_noms = ['Jan','Fév','Mar','Avr','Mai','Jun',
                 'Jul','Aoû','Sep','Oct','Nov','Déc']
    emprunts_par_mois = []
    labels_mois = []
    for i in range(11, -1, -1):
        date = aujourd_hui - timedelta(days=30*i)
        debut = date.replace(day=1)
        if debut.month == 12:
            fin = debut.replace(year=debut.year+1, month=1, day=1)
        else:
            fin = debut.replace(month=debut.month+1, day=1)
        count = Demande.objects.filter(
            date_demande__date__gte=debut,
            date_demande__date__lt=fin
        ).count()
        emprunts_par_mois.append(count)
        labels_mois.append(mois_noms[debut.month-1])

    # ── Graphique statuts demandes ──
    statuts_data = [
        demandes_en_attente,
        demandes_en_cours,
        demandes_restituees,
        demandes_refusees,
    ]

    # ── Graphique matériels par état ──
    etats_data = [
        materiels_disponibles,
        materiels_empruntes,
        materiels_maintenance,
        materiels_hors_service,
    ]

    # ── Graphique activités clubs par type ──
    types_activites = ['conference', 'sortie', 'jpo', 'atelier', 'cours', 'competition', 'autre']
    labels_activites = ['Conférence', 'Sortie', 'JPO', 'Atelier', 'Cours', 'Compétition', 'Autre']
    data_activites = [
        Activite.objects.filter(type=t).count() for t in types_activites
    ]

    # ── Graphique utilisateurs par rôle ──
    users_data = [
        total_etudiants,
        total_enseignants,
        total_techniciens,
        Utilisateur.objects.filter(role='admin').count(),
    ]

    contexte = {
        # Matériels
        'total_materiels': total_materiels,
        'materiels_disponibles': materiels_disponibles,
        'materiels_empruntes': materiels_empruntes,
        'materiels_maintenance': materiels_maintenance,
        'materiels_hors_service': materiels_hors_service,
        'taux_utilisation': taux_utilisation,
        # Demandes
        'total_demandes': total_demandes,
        'demandes_en_attente': demandes_en_attente,
        'demandes_en_cours': demandes_en_cours,
        'demandes_restituees': demandes_restituees,
        'demandes_refusees': demandes_refusees,
        'demandes_ce_mois': demandes_ce_mois,
        'demandes_cette_annee': demandes_cette_annee,
        'retards': retards,
        # Utilisateurs
        'total_etudiants': total_etudiants,
        'total_enseignants': total_enseignants,
        'total_techniciens': total_techniciens,
        'etudiants_bloques': etudiants_bloques,
        'total_penalites': total_penalites,
        # Clubs
        'total_clubs': total_clubs,
        'total_membres': total_membres,
        'total_activites': total_activites,
        'activites_ce_mois': activites_ce_mois,
        'activites_terminees': activites_terminees,
        # Maintenances
        'maintenances_en_cours': maintenances_en_cours,
        'maintenances_en_attente': maintenances_en_attente,
        # Listes
        'dernieres_demandes': dernieres_demandes,
        'top_materiels': top_materiels,
        'categories_stats': categories_stats,
        'clubs_stats': clubs_stats,
        'top_penalites': top_penalites,
        'top_etudiants': top_etudiants,
        # Graphiques
        'emprunts_par_mois': json.dumps(emprunts_par_mois),
        'labels_mois': json.dumps(labels_mois),
        'etats_data': json.dumps(etats_data),
        'statuts_data': json.dumps(statuts_data),
        'data_activites': json.dumps(data_activites),
        'labels_activites': json.dumps(labels_activites),
        'users_data': json.dumps(users_data),
    }
    return render(request, 'dashboard/dashboard.html', contexte)


@login_required
def export_pdf(request):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer, Image, HRFlowable
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    import io
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import matplotlib.patches as mpatches
    from django.db.models import Count

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4,
                           topMargin=2*cm, bottomMargin=2*cm,
                           leftMargin=2*cm, rightMargin=2*cm)
    elements = []
    styles = getSampleStyleSheet()

    titre_style = ParagraphStyle('titre', parent=styles['Title'],
                                 textColor=colors.HexColor('#1a5276'),
                                 fontSize=18, spaceAfter=6)
    h2_style = ParagraphStyle('h2', parent=styles['Heading2'],
                              textColor=colors.HexColor('#1a5276'), spaceAfter=8)
    sous_titre_style = ParagraphStyle('sous_titre', parent=styles['Normal'],
                                      textColor=colors.grey, fontSize=10,
                                      spaceAfter=20)

    def make_table(data, col_widths):
        t = Table(data, colWidths=col_widths)
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1a5276')),
            ('TEXTCOLOR', (0,0), (-1,0), colors.white),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
            ('FONTSIZE', (0,0), (-1,-1), 9),
            ('ALIGN', (0,0), (-1,-1), 'CENTER'),
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1),
             [colors.white, colors.HexColor('#f8fafc')]),
            ('TOPPADDING', (0,0), (-1,-1), 6),
            ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ]))
        return t

    def graphique_barres(labels, data, titre, couleur='#1a5276', largeur=15, hauteur=6):
        fig, ax = plt.subplots(figsize=(largeur/2.54, hauteur/2.54))
        barres = ax.bar(labels, data, color=couleur, alpha=0.85,
                       edgecolor='white', linewidth=0.5)
        ax.set_title(titre, fontsize=11, fontweight='bold', color='#1a5276', pad=10)
        ax.set_facecolor('#f8fafc')
        fig.patch.set_facecolor('white')
        ax.spines['top'].set_visible(False)
        ax.spines['right'].set_visible(False)
        ax.spines['left'].set_color('#e2e8f0')
        ax.spines['bottom'].set_color('#e2e8f0')
        ax.tick_params(axis='both', labelsize=8, colors='#64748b')
        ax.yaxis.set_tick_params(length=0)
        ax.set_axisbelow(True)
        ax.yaxis.grid(True, color='#f0f4f8', linewidth=0.8)
        for barre in barres:
            hauteur_barre = barre.get_height()
            if hauteur_barre > 0:
                ax.text(barre.get_x() + barre.get_width()/2., hauteur_barre + 0.1,
                       f'{int(hauteur_barre)}', ha='center', va='bottom',
                       fontsize=8, color='#1a5276', fontweight='bold')
        plt.tight_layout()
        img_buffer = io.BytesIO()
        plt.savefig(img_buffer, format='PNG', dpi=150, bbox_inches='tight')
        img_buffer.seek(0)
        plt.close()
        return Image(img_buffer, width=largeur*cm, height=hauteur*cm)

    def graphique_camembert(labels, data, titre, couleurs=None, largeur=8, hauteur=7):
        if couleurs is None:
            couleurs = ['#1a5276', '#27ae60', '#e67e22', '#e74c3c',
                       '#8e44ad', '#f39c12', '#95a5a6']
        data_nonzero = [(l, d, c) for l, d, c in zip(labels, data, couleurs) if d > 0]
        if not data_nonzero:
            return None
        labels_nz, data_nz, couleurs_nz = zip(*data_nonzero)
        fig, ax = plt.subplots(figsize=(largeur/2.54, hauteur/2.54))
        wedges, texts, autotexts = ax.pie(
            data_nz, labels=None, colors=couleurs_nz,
            autopct='%1.1f%%', startangle=90,
            wedgeprops={'edgecolor': 'white', 'linewidth': 2},
            pctdistance=0.75
        )
        for at in autotexts:
            at.set_fontsize(8)
            at.set_color('white')
            at.set_fontweight('bold')
        ax.set_title(titre, fontsize=11, fontweight='bold', color='#1a5276', pad=10)
        legend = ax.legend(
            wedges, [f'{l} ({d})' for l, d in zip(labels_nz, data_nz)],
            loc='lower center', bbox_to_anchor=(0.5, -0.25),
            ncol=2, fontsize=7, frameon=False
        )
        plt.tight_layout()
        img_buffer = io.BytesIO()
        plt.savefig(img_buffer, format='PNG', dpi=150, bbox_inches='tight')
        img_buffer.seek(0)
        plt.close()
        return Image(img_buffer, width=largeur*cm, height=hauteur*cm)

    aujourd_hui = timezone.now()
    mois_noms = ['Jan','Fév','Mar','Avr','Mai','Jun',
                 'Jul','Aoû','Sep','Oct','Nov','Déc']

    # ── En-tête ──
    elements.append(Paragraph("UFR Sciences de l'Ingénieur", titre_style))
    elements.append(Paragraph(
        f"Rapport complet de gestion — Généré le {aujourd_hui.strftime('%d/%m/%Y à %H:%M')}",
        sous_titre_style
    ))
    elements.append(HRFlowable(width="100%", thickness=2,
                               color=colors.HexColor('#1a5276')))
    elements.append(Spacer(1, 0.5*cm))

    # ── Stats globales ──
    elements.append(Paragraph("1. Statistiques globales", h2_style))
    data_stats = [
        ['Indicateur', 'Valeur'],
        ['Total matériels', str(Materiel.objects.count())],
        ['Matériels disponibles', str(Materiel.objects.filter(etat='disponible').count())],
        ['Matériels empruntés', str(Materiel.objects.filter(etat='emprunte').count())],
        ['En maintenance', str(Materiel.objects.filter(etat='maintenance').count())],
        ['Total emprunts', str(Demande.objects.count())],
        ['Emprunts en cours', str(Demande.objects.filter(statut='en_cours').count())],
        ['Emprunts en attente', str(Demande.objects.filter(statut='en_attente').count())],
        ['Total étudiants', str(Utilisateur.objects.filter(role='etudiant').count())],
        ['Étudiants bloqués', str(Utilisateur.objects.filter(bloque=True).count())],
        ['Total clubs', str(Club.objects.count())],
        ['Total activités clubs', str(Activite.objects.count())],
    ]
    elements.append(make_table(data_stats, [11*cm, 5*cm]))
    elements.append(Spacer(1, 0.5*cm))

    # ── Graphique emprunts par mois ──
    elements.append(Paragraph("2. Emprunts par mois (12 derniers mois)", h2_style))
    emprunts_par_mois = []
    labels_mois = []
    for i in range(11, -1, -1):
        date = aujourd_hui.date() - timedelta(days=30*i)
        debut = date.replace(day=1)
        if debut.month == 12:
            fin = debut.replace(year=debut.year+1, month=1, day=1)
        else:
            fin = debut.replace(month=debut.month+1, day=1)
        count = Demande.objects.filter(
            date_demande__date__gte=debut,
            date_demande__date__lt=fin
        ).count()
        emprunts_par_mois.append(count)
        labels_mois.append(mois_noms[debut.month-1])
    img = graphique_barres(labels_mois, emprunts_par_mois,
                           'Nombre d\'emprunts par mois', '#1a5276', 15, 7)
    elements.append(img)
    elements.append(Spacer(1, 0.5*cm))

    # ── Graphiques état matériels + statuts ──
    elements.append(Paragraph("3. Répartition matériels et emprunts", h2_style))
    from reportlab.platypus import HRFlowable
    from reportlab.platypus import KeepInFrame

    mat_labels = ['Disponible', 'Emprunté', 'Maintenance', 'Hors service']
    mat_data = [
        Materiel.objects.filter(etat='disponible').count(),
        Materiel.objects.filter(etat='emprunte').count(),
        Materiel.objects.filter(etat='maintenance').count(),
        Materiel.objects.filter(etat='hors_service').count(),
    ]
    mat_couleurs = ['#27ae60', '#e67e22', '#e74c3c', '#95a5a6']
    img_mat = graphique_camembert(mat_labels, mat_data, 'État du matériel',
                                  mat_couleurs, 8, 8)

    statuts_labels = ['En attente', 'En cours', 'Restituée', 'Refusée']
    statuts_data = [
        Demande.objects.filter(statut='en_attente').count(),
        Demande.objects.filter(statut='en_cours').count(),
        Demande.objects.filter(statut='restituee').count(),
        Demande.objects.filter(statut='refusee').count(),
    ]
    statuts_couleurs = ['#f39c12', '#1a5276', '#27ae60', '#e74c3c']
    img_statuts = graphique_camembert(statuts_labels, statuts_data,
                                      'Statuts des demandes',
                                      statuts_couleurs, 8, 8)

    if img_mat and img_statuts:
        t_imgs = Table([[img_mat, img_statuts]], colWidths=[8.5*cm, 8.5*cm])
        elements.append(t_imgs)
    elements.append(Spacer(1, 0.5*cm))

    # ── Top matériels ──
    elements.append(Paragraph("4. Top matériels les plus empruntés", h2_style))
    top_mat = Materiel.objects.annotate(nb=Count('lignes_demande')).order_by('-nb')[:10]
    top_labels = [m.nom[:15] for m in top_mat]
    top_data = [m.nb for m in top_mat]
    if any(d > 0 for d in top_data):
        img_top = graphique_barres(top_labels, top_data,
                                   'Top matériels empruntés', '#2980b9', 15, 7)
        elements.append(img_top)
    data_top_table = [['Matériel', 'N° Série', 'Catégorie', 'Nb emprunts']]
    for mat in top_mat:
        data_top_table.append([
            mat.nom, mat.numero_serie,
            str(mat.categorie) if mat.categorie else '—', str(mat.nb)
        ])
    elements.append(Spacer(1, 0.3*cm))
    elements.append(make_table(data_top_table, [6*cm, 4*cm, 4*cm, 3*cm]))
    elements.append(Spacer(1, 0.5*cm))

    # ── Clubs ──
    elements.append(Paragraph("5. Statistiques des clubs", h2_style))
    clubs_labels = []
    clubs_data = []
    data_clubs_table = [['Club', 'Membres', 'Activités', 'Score']]
    for club in Club.objects.annotate(
        nb_membres=Count('membres'),
        nb_activites=Count('activites')
    ):
        clubs_labels.append(club.nom[:12])
        clubs_data.append(club.nb_activites)
        data_clubs_table.append([
            club.nom, str(club.nb_membres),
            str(club.nb_activites), str(club.score_activite())
        ])
    if clubs_data and any(d > 0 for d in clubs_data):
        img_clubs = graphique_barres(clubs_labels, clubs_data,
                                     'Activités par club', '#8e44ad', 15, 6)
        elements.append(img_clubs)
    elements.append(Spacer(1, 0.3*cm))
    elements.append(make_table(data_clubs_table, [6*cm, 3*cm, 3*cm, 3*cm]))
    elements.append(Spacer(1, 0.5*cm))

    # ── Types activités ──
    elements.append(Paragraph("6. Types d'activités des clubs", h2_style))
    types_labels = ['Conférence', 'Sortie', 'JPO', 'Atelier', 'Cours', 'Compétition', 'Autre']
    types_ids = ['conference', 'sortie', 'jpo', 'atelier', 'cours', 'competition', 'autre']
    types_data = [Activite.objects.filter(type=t).count() for t in types_ids]
    img_types = graphique_barres(types_labels, types_data,
                                 'Types d\'activités', '#27ae60', 15, 6)
    elements.append(img_types)
    elements.append(Spacer(1, 0.5*cm))

    # ── Étudiants ──
    elements.append(Paragraph("7. Étudiants les plus actifs", h2_style))
    data_etu = [['Étudiant', 'Email', 'Nb emprunts', 'Pénalités', 'Score']]
    for etu in Utilisateur.objects.filter(role='etudiant').annotate(
        nb=Count('demandes')).order_by('-nb')[:10]:
        data_etu.append([
            etu.get_full_name() or etu.username,
            etu.email,
            str(etu.nb),
            str(etu.penalite),
            f"{etu.get_score_fiabilite()}/100"
        ])
    elements.append(make_table(data_etu, [5*cm, 4*cm, 3*cm, 2.5*cm, 2.5*cm]))
    elements.append(Spacer(1, 0.5*cm))

    # ── Dernières demandes ──
    elements.append(Paragraph("8. Dernières demandes (20)", h2_style))
    data_dem = [['N°', 'Étudiant', 'Matériel', 'Date', 'Statut']]
    for d in Demande.objects.all().order_by('-date_demande')[:20]:
        mat = ', '.join([l.materiel.nom for l in d.lignes.all()])
        data_dem.append([
            f'#{d.id}',
            (d.utilisateur.get_full_name() or d.utilisateur.username)[:20],
            mat[:20] + '...' if len(mat) > 20 else mat,
            d.date_demande.strftime('%d/%m/%Y'),
            d.statut
        ])
    elements.append(make_table(data_dem, [1.5*cm, 4.5*cm, 4*cm, 3*cm, 3*cm]))

    doc.build(elements)
    buffer.seek(0)
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = 'attachment; filename="rapport_complet_ufr_si.pdf"'
    return response


@login_required
def export_excel(request):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from django.db.models import Count
    import io

    wb = openpyxl.Workbook()
    couleur_header = 'FF1a5276'
    couleur_alt = 'FFf0f4f8'

    def style_header(cell):
        cell.font = Font(bold=True, color='FFFFFFFF', size=11)
        cell.fill = PatternFill(fill_type='solid', fgColor=couleur_header)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    def style_alt(cell, alt=False):
        if alt:
            cell.fill = PatternFill(fill_type='solid', fgColor=couleur_alt)
        cell.alignment = Alignment(horizontal='center', vertical='center')

    def auto_width(ws):
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_length + 4, 40)

    # Feuille 1 — Résumé
    ws0 = wb.active
    ws0.title = "Résumé"
    ws0.row_dimensions[1].height = 30
    headers0 = ['Indicateur', 'Valeur']
    for col, h in enumerate(headers0, 1):
        style_header(ws0.cell(row=1, column=col, value=h))
    data0 = [
        ('Total matériels', Materiel.objects.count()),
        ('Matériels disponibles', Materiel.objects.filter(etat='disponible').count()),
        ('Matériels empruntés', Materiel.objects.filter(etat='emprunte').count()),
        ('En maintenance', Materiel.objects.filter(etat='maintenance').count()),
        ('Total emprunts', Demande.objects.count()),
        ('En cours', Demande.objects.filter(statut='en_cours').count()),
        ('En attente', Demande.objects.filter(statut='en_attente').count()),
        ('Retards', Demande.objects.filter(statut='en_cours', date_fin__lt=timezone.now().date()).count()),
        ('Total étudiants', Utilisateur.objects.filter(role='etudiant').count()),
        ('Étudiants bloqués', Utilisateur.objects.filter(bloque=True).count()),
        ('Total clubs', Club.objects.count()),
        ('Total activités clubs', Activite.objects.count()),
    ]
    for row, (ind, val) in enumerate(data0, 2):
        ws0.cell(row=row, column=1, value=ind)
        ws0.cell(row=row, column=2, value=val)
        style_alt(ws0.cell(row=row, column=1), row % 2 == 0)
        style_alt(ws0.cell(row=row, column=2), row % 2 == 0)
    auto_width(ws0)

    # Feuille 2 — Matériels
    ws1 = wb.create_sheet("Matériels")
    headers1 = ['ID', 'Nom', 'Catégorie', 'N° Série', 'État', 'Date acquisition', 'Nb emprunts']
    for col, h in enumerate(headers1, 1):
        style_header(ws1.cell(row=1, column=col, value=h))
    for row, mat in enumerate(Materiel.objects.annotate(
        nb=Count('lignes_demande')).all(), 2):
        vals = [mat.id, mat.nom, str(mat.categorie) if mat.categorie else '',
                mat.numero_serie, mat.etat,
                str(mat.date_acquisition) if mat.date_acquisition else '', mat.nb]
        for col, v in enumerate(vals, 1):
            style_alt(ws1.cell(row=row, column=col, value=v), row % 2 == 0)
    auto_width(ws1)

    # Feuille 3 — Emprunts
    ws2 = wb.create_sheet("Emprunts")
    headers2 = ['ID', 'Étudiant', 'Email', 'Matériel', 'Date demande',
                'Date début', 'Date fin', 'Statut', 'Lieu']
    for col, h in enumerate(headers2, 1):
        style_header(ws2.cell(row=1, column=col, value=h))
    for row, dem in enumerate(Demande.objects.all().order_by('-date_demande'), 2):
        materiels = ', '.join([l.materiel.nom for l in dem.lignes.all()])
        lieu = dem.emplacement.libelle if hasattr(dem, 'emplacement') and dem.emplacement else ''
        vals = [
            dem.id,
            dem.utilisateur.get_full_name() or dem.utilisateur.username,
            dem.utilisateur.email,
            materiels,
            dem.date_demande.strftime('%d/%m/%Y %H:%M'),
            str(dem.date_debut),
            str(dem.date_fin),
            dem.statut,
            lieu
        ]
        for col, v in enumerate(vals, 1):
            style_alt(ws2.cell(row=row, column=col, value=v), row % 2 == 0)
    auto_width(ws2)

    # Feuille 4 — Étudiants
    ws3 = wb.create_sheet("Étudiants")
    headers3 = ['ID', 'Nom', 'Prénom', 'Email', 'Téléphone',
                'Pénalités', 'Bloqué', 'Nb emprunts', 'Score fiabilité']
    for col, h in enumerate(headers3, 1):
        style_header(ws3.cell(row=1, column=col, value=h))
    for row, etu in enumerate(Utilisateur.objects.filter(role='etudiant').annotate(
        nb=Count('demandes')), 2):
        vals = [
            etu.id, etu.last_name, etu.first_name, etu.email,
            etu.telephone, etu.penalite,
            'Oui' if etu.bloque else 'Non',
            etu.nb, etu.get_score_fiabilite()
        ]
        for col, v in enumerate(vals, 1):
            style_alt(ws3.cell(row=row, column=col, value=v), row % 2 == 0)
    auto_width(ws3)

    # Feuille 5 — Clubs
    ws4 = wb.create_sheet("Clubs")
    headers4 = ['Club', 'Membres', 'Activités totales',
                'Activités terminées', 'Cotisation (FCFA)']
    for col, h in enumerate(headers4, 1):
        style_header(ws4.cell(row=1, column=col, value=h))
    for row, club in enumerate(Club.objects.annotate(
        nb_membres=Count('membres'),
        nb_activites=Count('activites')
    ), 2):
        terminees = club.activites.filter(statut='terminee').count()
        vals = [club.nom, club.nb_membres, club.nb_activites,
                terminees, float(club.cotisation_montant)]
        for col, v in enumerate(vals, 1):
            style_alt(ws4.cell(row=row, column=col, value=v), row % 2 == 0)
    auto_width(ws4)

    # Feuille 6 — Activités
    ws5 = wb.create_sheet("Activités")
    headers5 = ['Club', 'Titre', 'Type', 'Date', 'Lieu', 'Statut']
    for col, h in enumerate(headers5, 1):
        style_header(ws5.cell(row=1, column=col, value=h))
    for row, act in enumerate(Activite.objects.all().order_by('-date'), 2):
        vals = [act.club.nom, act.titre, act.get_type_display(),
                str(act.date), act.lieu, act.get_statut_display()]
        for col, v in enumerate(vals, 1):
            style_alt(ws5.cell(row=row, column=col, value=v), row % 2 == 0)
    auto_width(ws5)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="rapport_complet_ufr_si.xlsx"'
    return response